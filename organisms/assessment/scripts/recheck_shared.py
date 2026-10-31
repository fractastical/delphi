#!/usr/bin/env python3
"""Re-check the shared trait/reference-database records of organisms/merged/datasets.json.

Batch: every merged record whose id does NOT start with one of the per-taxon query
prefixes (gbif-occurrence, ncbi-genome, ncbi-taxonomy, pbdb-, paleobiology-database-,
obis-occurrence, gbif-backbone, iucn-red-list-via-gbif). That leaves the shared databases
(AnAge, PanTHERIA, AVONET, Amniote DB, EltonTraits, FishBase, WoRMS, MorphoSource, IUCN,
Open Tree TNRS, TimeTree, Wikidata).

For each record the script
  * re-resolves `identifier`, `doi` and `is_referenced_by` (DOIs through
    https://doi.org/api/handles/<doi>, plain URLs through HTTP GET);
  * re-downloads every distribution that carries a recorded sha256 and recomputes it
    (files are cached under scripts/cache/, which is gitignored);
  * re-derives the record's counts from the downloaded file where the file is a table
    (AnAge, Amniote DB, PanTHERIA, EltonTraits, AVONET, FishBase parquet) and compares the
    recorded per-taxon row counts with rows matched on the taxon's accepted name at its
    rank (species: binomial; genus/family/order/class: the corresponding column);
  * re-runs the recorded API queries (WoRMS, MorphoSource, Wikidata, Open Tree TNRS,
    TimeTree, FishBase pages, IUCN) and compares the counts.

Output: scripts/recheck_shared_results.json, one object per record, in input order. Nothing
under organisms/<project>/ or organisms/merged/ is modified. The script never invents a
value: every observed number comes from a response or a downloaded file.

Usage: python3 recheck_shared.py [--ids id1,id2] [--offline]
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import io
import json
import re
import sys
import time
import zipfile
from pathlib import Path
from urllib.parse import unquote, urlparse

import pandas as pd
import requests

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
MERGED = ROOT / "organisms" / "merged" / "datasets.json"
TAXA = ROOT / "organisms" / "merged" / "taxa.json"
CACHE = HERE / "cache"
OUT = HERE / "recheck_shared_results.json"

EXCLUDED_PREFIXES = (
    "gbif-occurrence", "ncbi-genome", "ncbi-taxonomy", "pbdb-", "paleobiology-database-",
    "obis-occurrence", "gbif-backbone", "iucn-red-list-via-gbif",
)

UA = ("Mozilla/5.0 (X11; Linux x86_64) delphi-dataset-reliability-assessor/0.1 "
      "(+https://github.com/fractastical/delphi)")
TOL_REL = 0.005
TOL_ABS = 5

SESSION = requests.Session()
SESSION.headers["User-Agent"] = UA
OFFLINE = False


def now_iso() -> str:
    return dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def sha256_bytes(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


# --------------------------------------------------------------------------- HTTP + cache

def _cache_key(method: str, url: str, body: str | None) -> str:
    return hashlib.sha1(f"{method} {url}\n{body or ''}".encode()).hexdigest()


def fetch(url: str, method: str = "GET", body: str | None = None, headers: dict | None = None,
          timeout: int = 120, retries: int = 3, sleep: float = 0.0) -> dict:
    """Fetch with on-disk cache. Returns {status, content_type, bytes, from_cache, error}."""
    CACHE.mkdir(exist_ok=True)
    key = _cache_key(method, url, body)
    binf, metaf = CACHE / f"{key}.bin", CACHE / f"{key}.meta.json"
    if metaf.exists() and binf.exists():
        meta = json.loads(metaf.read_text())
        # transient failures are not worth caching
        if OFFLINE or meta.get("status") not in (429, 500, 502, 503, 504):
            meta["bytes"] = binf.read_bytes()
            meta["from_cache"] = True
            return meta
    if OFFLINE:
        return {"status": None, "content_type": None, "bytes": b"", "from_cache": False,
                "error": "offline and not cached"}
    err = None
    for attempt in range(retries):
        try:
            if sleep:
                time.sleep(sleep)
            h = dict(headers or {})
            if method == "POST":
                h.setdefault("Content-Type", "application/json")
                r = SESSION.post(url, data=body.encode() if body else None, headers=h, timeout=timeout)
            else:
                r = SESSION.get(url, headers=h, timeout=timeout, allow_redirects=True)
            if r.status_code in (429, 502, 503, 504) and attempt < retries - 1:
                wait = float(r.headers.get("Retry-After") or 0) or 5.0 * (attempt + 1)
                time.sleep(min(wait, 60))
                continue
            meta = {"status": r.status_code, "content_type": r.headers.get("Content-Type"),
                    "final_url": r.url, "fetched_at": now_iso(), "url": url, "method": method}
            binf.write_bytes(r.content)
            metaf.write_text(json.dumps(meta))
            meta["bytes"] = r.content
            meta["from_cache"] = False
            return meta
        except requests.RequestException as e:  # pragma: no cover - network
            err = str(e)
            time.sleep(2 * (attempt + 1))
    return {"status": None, "content_type": None, "bytes": b"", "from_cache": False, "error": err}


def resolve_doi(doi: str) -> dict:
    doi = doi.strip()
    doi = re.sub(r"^https?://(dx\.)?doi\.org/", "", doi)
    doi = re.sub(r"^doi:", "", doi, flags=re.I)
    r = fetch(f"https://doi.org/api/handles/{doi}")
    out = {"doi": doi, "handle_api_status": r["status"], "registered": None, "url": None}
    if r["status"] == 200:
        try:
            j = json.loads(r["bytes"])
            out["registered"] = j.get("responseCode") == 1
            for v in j.get("values", []):
                if v.get("type") == "URL":
                    out["url"] = v["data"]["value"]
        except json.JSONDecodeError:
            out["registered"] = False
    elif r["status"] == 404:
        out["registered"] = False
    return out


def resolve_url(url: str) -> dict:
    if url is None:
        return {"url": None, "status": None, "resolves": None}
    if re.match(r"^https?://(dx\.)?doi\.org/10\.", url):
        d = resolve_doi(url)
        return {"url": url, "via": "doi handle API", "status": d["handle_api_status"],
                "resolves": d["registered"], "target": d["url"]}
    r = fetch(url, timeout=60)
    status = r["status"]
    # figshare landing pages answer 202 to non-browser clients; esapubs directories answer 403
    resolves = status is not None and status < 400
    return {"url": url, "status": status, "resolves": resolves,
            "final_url": r.get("final_url"), "error": r.get("error")}


def within_tolerance(recorded, observed) -> dict:
    if recorded is None or observed is None:
        return {"recorded": recorded, "observed": observed, "within_tolerance": None}
    try:
        rec, obs = float(recorded), float(observed)
    except (TypeError, ValueError):
        return {"recorded": recorded, "observed": observed, "within_tolerance": recorded == observed}
    diff = abs(rec - obs)
    rel = diff / rec if rec else (0.0 if diff == 0 else None)
    ok = diff <= TOL_ABS or (rel is not None and rel <= TOL_REL)
    return {"recorded": recorded, "observed": observed, "abs_diff": diff,
            "relative_diff": None if rel is None else round(rel, 6), "within_tolerance": ok}


# --------------------------------------------------------------------------- taxa helpers

def load_taxa() -> dict:
    """Index merged taxa by slug and by merge key (gbif:…, wd:…, ncbi:…); dataset `taxa` lists use keys."""
    taxa = json.loads(TAXA.read_text())
    out = {}
    for t in taxa:
        out[t["slug"]] = t
        if t.get("key"):
            out[t["key"]] = t
    return out


def record_taxon_slugs(rec: dict, taxa: dict) -> list[str]:
    return [taxa[k]["slug"] if k in taxa else k for k in rec.get("taxa", [])]


def slug_for_name(name: str, rec: dict, taxa: dict) -> str | None:
    for k in rec.get("taxa", []):
        t = taxa.get(k)
        if t and t.get("accepted_name") == name:
            return t["slug"]
    return None


def taxon_name_rank(slug: str, taxa: dict) -> tuple[str | None, str | None]:
    t = taxa.get(slug)
    if not t:
        return None, None
    return t.get("accepted_name"), (t.get("rank") or "").lower()


def recorded_per_taxon(rec: dict) -> dict[str, int]:
    """Collect recorded per-taxon row counts from the several shapes slices used."""
    out: dict[str, int] = {}
    for key in ("rows_per_taxon", "rows_for_taxa", "taxa_coverage", "taxa_rows", "taxon_coverage"):
        v = rec.get(key)
        if isinstance(v, dict):
            for k, n in v.items():
                if isinstance(n, (int, float)):
                    out.setdefault(k, n)
    for c in rec.get("counts", []):
        pid, val, desc = c.get("property_id", ""), c.get("value"), c.get("description", "") or ""
        if not isinstance(val, (int, float)):
            continue
        m = re.match(r"^n_rows[:_](.+)$", pid)
        if m and m.group(1) not in ("for_taxa", "for_taxon", "for_species_rank_taxa"):
            out.setdefault(m.group(1), val)
            continue
        if c.get("taxon"):
            out.setdefault(c["taxon"], val)
            continue
        if pid in ("n_rows_for_taxon", "n_rows", "n_media", "n_media_for_taxon", "aphia_record"):
            m = (re.search(r"rows for (?:[A-Za-z .&]+\()?([a-z0-9-]+)\)?\s*$", desc)
                 or re.search(r"\bfor ([a-z0-9-]+)\s*$", desc)
                 or re.match(r"^([a-z0-9-]+):", desc))
            if m:
                out.setdefault(m.group(1), val)
            continue
        if pid.startswith("n_media_"):
            out.setdefault(pid[len("n_media_"):].replace("_", "-"), val)
            continue
        if pid == "n_physical_objects":
            m = re.search(r"\"([^\"]+)\"", desc)
            if m:
                out.setdefault(m.group(1).lower().replace(" ", "-"), val)
            continue
        if pid in ("n_rows_for_taxa",) and "=" in desc:
            for slug, n in re.findall(r"([a-z0-9-]+)=(\d+)", desc):
                out.setdefault(slug, int(n))
    return out


def file_level_recorded(rec: dict) -> dict[str, list]:
    """Recorded whole-file counts (n_species / n_rows / n_bird_species ...), by property id."""
    out: dict[str, list] = {}
    for c in rec.get("counts", []):
        pid = c.get("property_id")
        if pid in ("n_species", "n_rows", "n_bird_species", "n_mammal_species", "n_families",
                   "n_columns", "n_species_avonet1_birdlife", "n_shared_strings") and not c.get("taxon"):
            out.setdefault(pid, []).append({"value": c.get("value"), "slice": c.get("slice"),
                                            "description": c.get("description")})
    return out


def count_matches(df: pd.DataFrame, cols: dict[str, str], name: str, rank: str) -> int | None:
    """Rows matching a taxon at its rank. cols maps rank -> column name; 'binomial' -> column."""
    if rank in ("species", "subspecies"):
        parts = name.split()
        if rank == "subspecies" and len(parts) >= 3:
            name = " ".join(parts[:2])
        col = cols.get("binomial")
        if col is None and "genus" in cols and "species_epithet" in cols:
            g, s = name.split()[0], name.split()[1] if len(name.split()) > 1 else ""
            return int(((df[cols["genus"]] == g) & (df[cols["species_epithet"]] == s)).sum())
        if col is None:
            return None
        return int((df[col].astype(str).str.strip() == name).sum())
    col = cols.get(rank)
    if col is None:
        return None
    return int((df[col].astype(str).str.strip() == name).sum())


def compare_taxa(rec: dict, df: pd.DataFrame, cols: dict[str, str], taxa: dict) -> dict:
    recorded = recorded_per_taxon(rec)
    comps = []
    for slug, n in recorded.items():
        name, rank = taxon_name_rank(slug, taxa)
        if not name:
            comps.append({"taxon": slug, "recorded": n, "observed": None, "note": "slug not in merged taxa"})
            continue
        obs = count_matches(df, cols, name, rank)
        entry = {"taxon": slug, "rank": rank, "name": name, "recorded": n, "observed": obs,
                 "match": None if obs is None else (obs == n)}
        if entry["match"] is False and rank in ("species", "subspecies") and cols.get("binomial"):
            # the vendor may have matched a synonym (e.g. Neovison/Mustela vison for Neogale vison)
            epithet = name.split()[1] if len(name.split()) > 1 else None
            if epithet:
                col = df[cols["binomial"]].astype(str)
                hits = col[col.str.split().str[1] == epithet].unique().tolist()
                entry["binomials_sharing_epithet"] = hits[:6]
        comps.append(entry)
    n_cmp = sum(1 for c in comps if c.get("observed") is not None)
    n_ok = sum(1 for c in comps if c.get("match"))
    # species-rank taxa listed on the record but absent from the table
    missing = []
    for slug in record_taxon_slugs(rec, taxa):
        name, rank = taxon_name_rank(slug, taxa)
        if name and rank in ("species",):
            obs = count_matches(df, cols, name, rank)
            if obs == 0:
                missing.append(slug)
    return {"n_recorded_per_taxon": len(recorded), "n_compared": n_cmp, "n_match": n_ok,
            "mismatches": [c for c in comps if c.get("match") is False][:60],
            "unmatched_species_rank_taxa": missing}


# --------------------------------------------------------------------------- per-source file parsers

def parse_amniote(b: bytes) -> tuple[pd.DataFrame, dict]:
    df = pd.read_csv(io.BytesIO(b), low_memory=False)
    info = {"n_rows": int(len(df)), "n_columns": int(df.shape[1]), "columns_head": list(df.columns[:8])}
    cols = {"class": "class", "order": "order", "family": "family", "genus": "genus",
            "species_epithet": "species"}
    df["_binomial"] = df["genus"].astype(str).str.strip() + " " + df["species"].astype(str).str.strip()
    cols["binomial"] = "_binomial"
    return df, {"info": info, "cols": cols}


def parse_anage_zip(b: bytes) -> tuple[pd.DataFrame, dict]:
    z = zipfile.ZipFile(io.BytesIO(b))
    names = z.namelist()
    data = z.read("anage_data.txt")
    df = pd.read_csv(io.BytesIO(data), sep="\t", low_memory=False, encoding="latin-1")
    release = ""
    for n in names:
        if n.lower().endswith(".html"):
            release = z.read(n).decode("utf-8", "replace")
    text = re.sub(r"<[^>]+>", " ", release)
    text = re.sub(r"\s+", " ", text)
    m_build = re.search(r"[Bb]uild\s*(\d+)", text)
    m_date = re.search(r"Release date:\s*([A-Za-z]+\s+\d{1,2},?\s+\d{4})", text)
    m_entries = re.search(r"Number of entries:\s*([\d,]+)", text)
    cites = re.findall(r"((?:Tacutu|de Magalh[^ ]+)[^.]*\(\d{4}\)[^.]*\.[^.]*\.)", text)
    dois = re.findall(r"10\.\d{4,9}/[^\s\"<>]+", release)
    info = {"zip_members": names, "n_rows": int(len(df)), "n_columns": int(df.shape[1]),
            "release_build": m_build.group(1) if m_build else None,
            "release_date_text": m_date.group(1) if m_date else None,
            "release_entries_statement": m_entries.group(1) if m_entries else None,
            "release_citation_snippets": cites[:3],
            "release_dois": sorted(set(dois)),
            "release_text_head": text[:1500]}
    df["_binomial"] = df["Genus"].astype(str).str.strip() + " " + df["Species"].astype(str).str.strip()
    cols = {"phylum": "Phylum", "class": "Class", "order": "Order", "family": "Family",
            "genus": "Genus", "binomial": "_binomial"}
    return df, {"info": info, "cols": cols}


def parse_pantheria(b: bytes) -> tuple[pd.DataFrame, dict]:
    df = pd.read_csv(io.BytesIO(b), sep="\t", low_memory=False, encoding="latin-1")
    info = {"n_rows": int(len(df)), "n_columns": int(df.shape[1]),
            "n_rows_with_binomial": int(df["MSW05_Binomial"].notna().sum())}
    cols = {"order": "MSW05_Order", "family": "MSW05_Family", "genus": "MSW05_Genus",
            "binomial": "MSW05_Binomial"}
    df["_class"] = "Mammalia"
    cols["class"] = "_class"
    return df, {"info": info, "cols": cols}


def parse_elton(b: bytes, kind: str) -> tuple[pd.DataFrame, dict]:
    text = b.decode("latin-1")
    lines = text.split("\n")
    df = pd.read_csv(io.StringIO(text), sep="\t", low_memory=False)
    name_col = "Scientific"
    info = {"n_lines_total": len(lines), "n_lines_nonempty": sum(1 for l in lines if l.strip()),
            "n_rows_pandas": int(len(df)),
            "n_rows_with_scientific_name": int(df[name_col].notna().sum()) if name_col in df else None,
            "n_columns": int(df.shape[1])}
    # The class column is only set on rows that carry a scientific name, so that
    # trailing citation/blank lines (which pandas still parses as rows) are not counted.
    has_name = df[name_col].notna() if name_col in df else pd.Series(False, index=df.index)
    if kind == "bird":
        info["n_rows_with_specid"] = int(pd.to_numeric(df["SpecID"], errors="coerce").notna().sum())
        cols = {"order": "IOCOrder", "family": "BLFamilyLatin", "binomial": "Scientific"}
        df["_class"] = has_name.map({True: "Aves", False: ""})
    else:
        info["n_rows_with_msw3_id"] = int(pd.to_numeric(df["MSW3_ID"], errors="coerce").notna().sum())
        cols = {"family": "MSWFamilyLatin", "binomial": "Scientific"}
        df["_class"] = has_name.map({True: "Mammalia", False: ""})
    df["_genus"] = df[name_col].astype(str).str.split().str[0]
    cols["genus"] = "_genus"
    cols["class"] = "_class"
    return df, {"info": info, "cols": cols}


def parse_avonet(b: bytes) -> tuple[pd.DataFrame, dict]:
    import openpyxl  # noqa: WPS433
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    sheets = wb.sheetnames
    ws = wb["AVONET1_BirdLife"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    data = [r for r in rows if r and r[0] is not None]
    df = pd.DataFrame(data, columns=header[: len(data[0])] if data else header)
    info = {"sheets": sheets, "n_rows_avonet1_birdlife": int(len(df)), "n_columns": len(header)}
    cols = {"order": "Order1", "family": "Family1", "binomial": "Species1"}
    df["_genus"] = df["Species1"].astype(str).str.split().str[0]
    df["_class"] = "Aves"
    cols["genus"] = "_genus"
    cols["class"] = "_class"
    return df, {"info": info, "cols": cols}


def parse_fishbase_parquet(species_b: bytes, families_b: bytes) -> tuple[pd.DataFrame, dict]:
    sp = pd.read_parquet(io.BytesIO(species_b))
    fam = pd.read_parquet(io.BytesIO(families_b))
    info = {"n_species_rows": int(len(sp)), "n_family_rows": int(len(fam)),
            "species_columns_head": list(sp.columns[:8])}
    df = sp.merge(fam[["FamCode", "Family", "Order", "Class"]], on="FamCode", how="left", suffixes=("", "_fam"))
    df["_binomial"] = df["Genus"].astype(str).str.strip() + " " + df["Species"].astype(str).str.strip()
    cols = {"class": "Class", "order": "Order", "family": "Family", "genus": "Genus", "binomial": "_binomial"}
    return df, {"info": info, "cols": cols}


# --------------------------------------------------------------------------- distribution checks

LIVE_HOSTS = ("www.fishbase.se", "www.wikidata.org", "www.morphosource.org", "www.marinespecies.org",
              "query.wikidata.org", "api.opentreeoflife.org", "timetree.org")


def check_distribution(d: dict, rec_id: str) -> dict:
    url = d["download_url"]
    out = {"download_url": url, "recorded_sha256": d.get("sha256"), "recorded_byte_size": d.get("byte_size")}
    host = urlparse(url).netloc
    method, body = "GET", None
    if "opentreeoflife.org/v3/tnrs/match_names" in url:
        method = "POST"
        body = d.get("request_body")
        if body is None and isinstance(d.get("request"), dict):
            body = json.dumps(d["request"])
        if body is None:
            m = re.search(r"(\{\"names\".*\})", d.get("description", "") or "")
            body = m.group(1) if m else None
        if body is not None:
            try:
                json.loads(body)
            except json.JSONDecodeError:
                body = None  # e.g. the placeholder "[<58 accepted names>]"
        if body is None:
            out.update({"status": None, "observed_sha256": None, "match": None,
                        "verdict": "not-reproducible", "note": "exact POST body not recorded on the distribution"})
            return out
        out["replayed_body_bytes"] = len(body.encode())
    if "iucnredlist.org" in host:
        r = fetch(url, timeout=30, retries=1)
        out.update({"status": r["status"], "observed_sha256": None, "match": None,
                    "verdict": "blocked" if r["status"] in (401, 403) else "no-checksum-recorded",
                    "note": f"IUCN answers HTTP {r['status']} to non-browser clients"})
        return out
    r = fetch(url, method=method, body=body, headers={"Accept": "application/json"} if "morphosource" in host or "opentree" in host else None,
              sleep=2.0 if "morphosource" in host else (0.3 if ("marinespecies" in host or "wikidata" in host) else 0.0),
              retries=4 if "morphosource" in host else 3)
    out["status"] = r["status"]
    b = r["bytes"]
    out["observed_byte_size"] = len(b)
    if r["status"] is None:
        out.update({"observed_sha256": None, "match": None, "verdict": "fetch-error", "note": r.get("error")})
        return out
    if r["status"] >= 400:
        out.update({"observed_sha256": None, "match": None,
                    "verdict": "blocked" if r["status"] in (401, 403, 429) else "http-error"})
        return out
    sha = sha256_bytes(b)
    out["observed_sha256"] = sha
    if d.get("sha256"):
        out["match"] = sha == d["sha256"]
        out["verdict"] = "match" if out["match"] else "mismatch"
        if not out["match"]:
            note = d.get("note") or d.get("description") or ""
            if "keys sorted" in note or "reduced to" in note or "normalized" in note.lower():
                # derived checksum: try the recorded derivation
                alt = derived_checksums(b, note)
                if d["sha256"] in alt.values():
                    out["match"] = True
                    out["verdict"] = "match"
                    out["note"] = "matches recorded derivation: " + [k for k, v in alt.items() if v == d["sha256"]][0]
                else:
                    live = host in LIVE_HOSTS
                    out["verdict"] = "mismatch-live" if live else "mismatch-derived"
                    out["note"] = ("recorded checksum is of a derived form (" + note[:80] + "); none of the tried derivations reproduce it"
                                   + ("; live API response, so a changed count/facet is the expected cause" if live else ""))
            elif host in LIVE_HOSTS:
                out["verdict"] = "mismatch-live"
                out["note"] = "live page/API response; byte-identity is not expected to persist"
    else:
        out["match"] = None
        out["verdict"] = "no-checksum-recorded"
    if d.get("byte_size") is not None:
        out["byte_size_match"] = (d["byte_size"] == len(b))
    return out


def derived_checksums(b: bytes, note: str) -> dict[str, str]:
    out = {}
    try:
        j = json.loads(b)
    except json.JSONDecodeError:
        return out
    for name, s in {
        "json sort_keys default separators": json.dumps(j, sort_keys=True),
        "json sort_keys compact": json.dumps(j, sort_keys=True, separators=(",", ":")),
        "json sort_keys ensure_ascii=False": json.dumps(j, sort_keys=True, ensure_ascii=False),
        "json sort_keys indent=2": json.dumps(j, sort_keys=True, indent=2),
    }.items():
        out[name] = sha256_bytes(s.encode())
    # MorphoSource wraps everything in a top-level "response" object; the vendor's
    # "reduced" checksum keeps that wrapper and filters the facet list by name.
    inner = j.get("response", j) if isinstance(j, dict) else None
    if isinstance(inner, dict) and "pages" in inner:
        facets = inner.get("facets") or []
        keep = ("media_type", "modality", "publication_status")
        if isinstance(facets, list):
            kept = [f for f in facets if isinstance(f, dict) and f.get("name") in keep]
        else:
            kept = {k: v for k, v in facets.items() if any(t in k for t in keep)}
        core = {"pages": inner.get("pages"), "facets": kept}
        for label, reduced in (("wrapped", {"response": core}), ("unwrapped", core)):
            for name, s in {
                f"reduced response(pages+facets) {label} sort_keys": json.dumps(reduced, sort_keys=True),
                f"reduced response(pages+facets) {label} compact": json.dumps(reduced, sort_keys=True, separators=(",", ":")),
                f"reduced response(pages+facets) {label} default": json.dumps(reduced),
            }.items():
                out[name] = sha256_bytes(s.encode())
    return out


# --------------------------------------------------------------------------- API re-queries

def morphosource_total(b: bytes) -> int | None:
    try:
        j = json.loads(b)
        j = j.get("response", j)
        return int(j["pages"]["total_count"])
    except Exception:  # noqa: BLE001
        return None


def recheck_morphosource(rec: dict, dist_results: list[dict], taxa: dict) -> dict:
    recorded = recorded_per_taxon(rec)
    comps = []
    for d, res in zip(rec["distributions"], dist_results):
        url = d["download_url"]
        if res.get("status") != 200:
            continue
        key = _cache_key("GET", url, None)
        b = (CACHE / f"{key}.bin").read_bytes()
        total = morphosource_total(b)
        slug = d.get("taxon")
        if not slug:
            q = re.search(r"[?&]q=([^&]+)", url) or re.search(r"taxonomy_gbif%5D%5B%5D=([^&]+)", url)
            qname = unquote(q.group(1)).strip('"') if q else None
            slug = slug_for_name(qname, rec, taxa) if qname else None
            if slug is None and qname:
                slug = qname.lower().replace(" ", "-")
            if slug not in recorded and qname:
                alt = qname.lower().replace(" ", "-")
                if alt in recorded:
                    slug = alt
        recn = recorded.get(slug) if slug else None
        comps.append({"taxon": slug, "url": url, **within_tolerance(recn, total)})
    n_cmp = sum(1 for c in comps if c.get("recorded") is not None and c.get("observed") is not None)
    n_ok = sum(1 for c in comps if c.get("within_tolerance"))
    return {"n_queries_rerun": len(comps), "n_compared": n_cmp, "n_within_tolerance": n_ok,
            "mismatches": [c for c in comps if c.get("within_tolerance") is False],
            "unmapped": [c for c in comps if c.get("recorded") is None][:10]}


def recheck_worms(rec: dict, dist_results: list[dict]) -> dict:
    comps = []
    n_no_content = 0
    for d, res in zip(rec["distributions"], dist_results):
        url = d["download_url"]
        if res.get("status") == 204:
            # WoRMS answers 204 for names it does not hold (terrestrial taxa); the slices document this
            n_no_content += 1
            comps.append({"url": url, "status": 204, "ok": True, "note": "no content: name not in WoRMS"})
            continue
        if res.get("status") not in (200, 206):
            comps.append({"url": url, "status": res.get("status"), "ok": False})
            continue
        key = _cache_key("GET", url, None)
        b = (CACHE / f"{key}.bin").read_bytes()
        entry = {"url": url, "status": res.get("status")}
        if "AphiaRecordByAphiaID" in url:
            j = json.loads(b)
            entry.update({"aphia_id": j.get("AphiaID"), "scientificname": j.get("scientificname"),
                          "status_worms": j.get("status"), "rank": j.get("rank")})
            desc = d.get("description", "") or ""
            m = re.match(r"Aphia record (\d+): (.*?) \((\w+), (\w+)", desc)
            if m:
                sci = j.get("scientificname") or ""
                entry["name_matches_description"] = m.group(2).startswith(sci) and m.group(3) == j.get("rank")
            else:
                entry["name_matches_description"] = True
            entry["ok"] = str(j.get("AphiaID")) == url.rstrip("/").split("/")[-1] and entry["name_matches_description"]
            if d.get("aphia_id") is not None:
                entry["ok"] = entry["ok"] and (j.get("AphiaID") == d["aphia_id"])
        elif "AphiaRecordsByName" in url:
            try:
                j = json.loads(b)
            except json.JSONDecodeError:
                j = []
            entry.update({"n_records": len(j) if isinstance(j, list) else None,
                          "names": sorted({x.get("scientificname") for x in j})[:5] if isinstance(j, list) else None})
            entry["ok"] = isinstance(j, list) and len(j) > 0
        elif "AphiaIDByName" in url:
            txt = b.decode("utf-8", "replace").strip()
            entry.update({"aphia_id_text": txt[:20]})
            entry["ok"] = bool(re.match(r"^-?\d+$", txt))
        else:
            entry["ok"] = True
        comps.append(entry)
    n_ok = sum(1 for c in comps if c.get("ok"))
    out = {"n_queries_rerun": len(comps), "n_ok": n_ok, "n_no_content_204": n_no_content,
           "failures": [c for c in comps if not c.get("ok")][:40]}
    # Records that vendored only an example query but claim "n_taxa_resolved": re-run
    # AphiaRecordsByName for every listed taxon and count the names WoRMS holds.
    recorded_resolved = next((c["value"] for c in rec.get("counts", []) or []
                              if isinstance(c, dict) and c.get("property_id") == "n_taxa_resolved"), None)
    if recorded_resolved is not None and len(rec["distributions"]) < len(rec.get("taxa", [])):
        taxa = load_taxa()
        resolved, unresolved, errors = 0, [], 0
        for slug in record_taxon_slugs(rec, taxa):
            name, _rank = taxon_name_rank(slug, taxa)
            if not name:
                continue
            url = ("https://www.marinespecies.org/rest/AphiaRecordsByName/"
                   + requests.utils.quote(name) + "?like=false&marine_only=false")
            r = fetch(url, sleep=0.3, retries=3)
            if r["status"] == 200:
                try:
                    j = json.loads(r["bytes"])
                except json.JSONDecodeError:
                    j = []
                if isinstance(j, list) and j:
                    resolved += 1
                else:
                    unresolved.append(name)
            elif r["status"] == 204:
                unresolved.append(name)
            else:
                errors += 1
        out["taxa_resolution"] = {"recorded_n_taxa_resolved": recorded_resolved, "n_taxa_in_record": len(rec.get("taxa", [])),
                                  "observed_names_with_aphia_records": resolved, "n_http_errors": errors,
                                  "unresolved_names": unresolved[:60],
                                  **within_tolerance(recorded_resolved, resolved)}
    return out


def _p225_values(entity: dict) -> set[str]:
    vals = set()
    for cl in (entity.get("claims") or {}).get("P225", []):
        try:
            vals.add(cl["mainsnak"]["datavalue"]["value"])
        except KeyError:
            pass
    return vals


def recheck_wikidata(rec: dict, dist_results: list[dict], taxa: dict) -> dict:
    comps = []
    n_entities = n_with_p225 = 0
    sparql_rows = None
    searched_names: list[str] = []
    all_p225: dict[str, set[str]] = {}
    single_entity_checks = []
    for d, res in zip(rec["distributions"], dist_results):
        url = d["download_url"]
        if res.get("status") != 200:
            comps.append({"url": url, "status": res.get("status"), "ok": False})
            continue
        key = _cache_key("GET", url, None)
        b = (CACHE / f"{key}.bin").read_bytes()
        j = json.loads(b)
        if "wbgetentities" in url:
            ents = j.get("entities", {})
            for qid, e in ents.items():
                n_entities += 1
                vals = _p225_values(e)
                all_p225[qid] = vals
                if vals:
                    n_with_p225 += 1
            # single-entity fetches "Qnnn entity (P225 checked) for <slug>": compare P225 with the accepted name
            m = re.search(r"\bfor ([a-z0-9-]+)\s*$", d.get("description", "") or "")
            if m and len(ents) == 1:
                slug = m.group(1)
                name, _ = taxon_name_rank(slug, taxa)
                qid = next(iter(ents))
                ok = name in all_p225[qid] if name else None
                single_entity_checks.append({"qid": qid, "taxon": slug, "accepted_name": name,
                                             "p225": sorted(all_p225[qid]), "ok": ok})
        elif "wbsearchentities" in url:
            m = re.search(r"[?&]search=([^&]+)", url)
            if m:
                searched_names.append(unquote(m.group(1)))
            if not j.get("search"):
                comps.append({"url": url, "ok": False, "note": "no search hits"})
        elif "sparql" in url:
            sparql_rows = len(j.get("results", {}).get("bindings", []))
    out = {"n_entities_fetched": n_entities, "n_entities_with_p225": n_with_p225, "failures": comps[:20]}
    if single_entity_checks:
        out["p225_equals_accepted_name"] = {
            "n_checked": len(single_entity_checks),
            "n_ok": sum(1 for c in single_entity_checks if c["ok"]),
            "failures": [c for c in single_entity_checks if not c["ok"]][:20]}
    if searched_names:
        # a searched name is resolved when some fetched entity carries it as P225
        union = set().union(*all_p225.values()) if all_p225 else set()
        resolved = [n for n in searched_names if n in union]
        rec_res = next((c["value"] for c in rec.get("counts", []) if c["property_id"] == "n_taxa_resolved"), None)
        out["searched_names_resolved"] = {**within_tolerance(rec_res, len(resolved)),
                                          "n_searched": len(searched_names),
                                          "unresolved": [n for n in searched_names if n not in union]}
    if sparql_rows is not None:
        rec_rows = next((c["value"] for c in rec.get("counts", []) if c["property_id"] == "n_rows"), None)
        out["sparql_rows"] = within_tolerance(rec_rows, sparql_rows)
    return out


def recheck_opentree(rec: dict, dist_results: list[dict], taxa: dict) -> dict:
    out: dict = {}
    about = fetch("https://api.opentreeoflife.org/v3/taxonomy/about", method="POST", body="{}")
    if about["status"] == 200:
        try:
            out["taxonomy_about"] = {k: v for k, v in json.loads(about["bytes"]).items() if k in ("source", "version", "name")}
        except json.JSONDecodeError:
            pass
    d, res = rec["distributions"][0], dist_results[0]
    body = None
    if res.get("status") == 200 and res.get("replayed_body_bytes"):
        body = d.get("request_body") or (json.dumps(d["request"]) if isinstance(d.get("request"), dict) else None)
        if body is None:
            m = re.search(r"(\{\"names\".*\})", d.get("description", "") or "")
            body = m.group(1) if m else None
        how = "recorded request body replayed"
    if body is None:
        names = []
        for s in record_taxon_slugs(rec, taxa):
            n, _ = taxon_name_rank(s, taxa)
            if n:
                names.append(n)
        body = json.dumps({"names": names, "context_name": "Animals", "do_approximate_matching": False})
        how = "request body not recorded; rebuilt from the record's taxa accepted names"
    r = fetch("https://api.opentreeoflife.org/v3/tnrs/match_names", method="POST", body=body)
    out["replay"] = {"how": how, "status": r["status"]}
    if r["status"] == 200:
        j = json.loads(r["bytes"])
        results = j.get("results", [])
        n_q = len(json.loads(body)["names"])
        n_matched = sum(1 for x in results if any(not m.get("is_approximate_match") for m in x.get("matches", [])))
        n_ambig = sum(1 for x in results if len([m for m in x.get("matches", []) if not m.get("is_approximate_match")]) > 1)
        unmatched = j.get("unmatched_names", [])
        out["replay"].update({"n_names_queried": n_q, "n_names_matched": n_matched,
                              "n_names_ambiguous": n_ambig, "unmatched_names": unmatched[:30],
                              "taxonomy": j.get("taxonomy")})
        rec_counts = {c["property_id"]: c["value"] for c in rec.get("counts", [])}
        for pid, obs in (("n_names_queried", n_q), ("n_names_matched", n_matched), ("n_taxa_matched", n_matched),
                         ("n_names_ambiguous", n_ambig), ("n_names_unmatched", len(unmatched))):
            if pid in rec_counts:
                out.setdefault("count_rechecks", {})[pid] = within_tolerance(rec_counts[pid], obs)
    return out


def recheck_timetree(rec: dict, dist_results: list[dict]) -> dict:
    comps = []
    rec_counts = {c["property_id"]: c for c in rec.get("counts", [])}
    for d, res in zip(rec["distributions"], dist_results):
        if res.get("status") != 200:
            comps.append({"url": d["download_url"], "status": res.get("status"), "ok": False})
            continue
        key = _cache_key("GET", d["download_url"], None)
        txt = (CACHE / f"{key}.bin").read_bytes().decode("utf-8", "replace")
        row = None
        try:
            df = pd.read_csv(io.StringIO(txt))
            row = df.iloc[0].to_dict() if len(df) else None
        except Exception:  # noqa: BLE001
            pass
        desc = d.get("description", "")
        pair = "__".join(re.sub(r"[^a-z0-9]+", "_", p.strip().lower()).strip("_") for p in desc.split(" vs "))
        entry = {"pair": desc, "raw_head": txt[:200]}
        if row:
            age = row.get("precomputed_age") or row.get("adjusted_age") or row.get("median_time")
            total = row.get("all_total") or row.get("total")
            entry.update({"precomputed_age": age, "all_total": total})
            key_age = f"divergence_mya_{pair}"
            key_n = f"n_studies_{pair}"
            if key_age in rec_counts:
                entry["age_check"] = within_tolerance(rec_counts[key_age]["value"], age)
            if key_n in rec_counts:
                entry["n_studies_check"] = within_tolerance(rec_counts[key_n]["value"], total)
        comps.append(entry)
    return {"pairs": comps}


def recheck_fishbase_pages(rec: dict, dist_results: list[dict]) -> dict:
    comps = []
    for d, res in zip(rec["distributions"], dist_results):
        url = d["download_url"]
        entry = {"url": url, "status": res.get("status")}
        if res.get("status") == 200:
            key = _cache_key("GET", url, None)
            html = (CACHE / f"{key}.bin").read_bytes().decode("utf-8", "replace")
            if "PopGrowthList" in url:
                # the growth-parameter table is the largest <table>; its first row is the header
                tables = re.findall(r"<table[^>]*>(.*?)</table>", html, flags=re.S | re.I)
                main = max(tables, key=len) if tables else ""
                rows = re.findall(r"<tr[^>]*>(.*?)</tr>", main, flags=re.S | re.I)
                entry["n_growth_rows"] = max(len(rows) - 1, 0)
                m = re.search(r"\[n = (\d+)\]", html)
                entry["lm_vs_linf_graph_n"] = int(m.group(1)) if m else None
            else:
                m = re.search(r"Max length\s*:?\s*</span>?\s*([\d.,]+\s*cm\s*[A-Z]{1,3})", html) or \
                    re.search(r"Max length\s*:\s*([\d.,]+\s*cm\s*[A-Z]{1,3})", re.sub(r"<[^>]+>", "", html))
                entry["max_length_text"] = m.group(1) if m else None
                mt = re.search(r"<title>(.*?)</title>", html, flags=re.S | re.I)
                entry["title"] = re.sub(r"\s+", " ", mt.group(1)).strip()[:80] if mt else None
        comps.append(entry)
    rec_growth = next((c["value"] for c in rec.get("counts", []) if c["property_id"] == "n_growth_studies"), None)
    out = {"pages": comps}
    if rec_growth is not None:
        obs = next((c.get("n_growth_rows") for c in comps if "PopGrowthList" in c["url"]), None)
        out["n_growth_studies"] = within_tolerance(rec_growth, obs)
    return out


# --------------------------------------------------------------------------- figshare metadata

# figshare file id -> article id, for files recorded only by their ndownloader URL
FIGSHARE_FILE_TO_ARTICLE = {"8067269": "3563457", "5631081": "3559887", "5631084": "3559887",
                            "34480856": "16586228"}


def figshare_articles(rec: dict) -> list[dict]:
    ids: set[str] = set()
    urls = [rec.get("identifier") or ""] + list(rec.get("alternate_identifiers", []))
    urls += [d["download_url"] for d in rec.get("distributions", [])]
    for u in urls:
        m = re.search(r"figshare\.com/articles/[^/]+/[^/]+/(\d+)", u) or re.search(r"api\.figshare\.com/v2/articles/(\d+)", u) \
            or re.search(r"m9\.figshare\.(\d+)", u)
        if m:
            ids.add(m.group(1))
        m = re.search(r"ndownloader\.figshare\.com/files/(\d+)", u)
        if m and m.group(1) in FIGSHARE_FILE_TO_ARTICLE:
            ids.add(FIGSHARE_FILE_TO_ARTICLE[m.group(1)])
    out = []
    for aid in sorted(ids):
        fa = fetch(f"https://api.figshare.com/v2/articles/{aid}")
        if fa["status"] != 200:
            out.append({"article_id": aid, "status": fa["status"]})
            continue
        j = json.loads(fa["bytes"])
        out.append({"article_id": aid, "title": j.get("title"), "version": j.get("version"), "doi": j.get("doi"),
                    "license": (j.get("license") or {}).get("name"),
                    "published_date": j.get("published_date"), "modified_date": j.get("modified_date"),
                    "files": [{"id": f["id"], "name": f["name"], "size": f["size"], "md5": f.get("computed_md5")}
                              for f in j.get("files", [])][:12]})
    return out


# --------------------------------------------------------------------------- main per-record

def cached_bytes(url: str) -> bytes | None:
    key = _cache_key("GET", url, None)
    f = CACHE / f"{key}.bin"
    meta = CACHE / f"{key}.meta.json"
    if f.exists() and meta.exists() and json.loads(meta.read_text()).get("status") == 200:
        return f.read_bytes()
    return None


def recheck_record(rec: dict, taxa: dict) -> dict:
    rid = rec["id"]
    print(f"[{now_iso()}] {rid}", file=sys.stderr, flush=True)
    out = {"dataset_id": rid, "checked_at": now_iso(), "slices": rec.get("slices", []),
           "n_taxa": len(rec.get("taxa", []))}
    out["identifier"] = resolve_url(rec.get("identifier"))
    out["doi"] = resolve_doi(rec["doi"]) if rec.get("doi") else None
    irb = rec.get("is_referenced_by")
    out["is_referenced_by"] = resolve_url(irb) if irb else None
    out["alternate_identifiers"] = [resolve_url(u) for u in rec.get("alternate_identifiers", [])]

    dist_results = [check_distribution(d, rid) for d in rec.get("distributions", [])]
    out["distributions"] = dist_results
    verdicts = [d.get("verdict") for d in dist_results]
    out["checksum_summary"] = {
        "n_distributions": len(dist_results),
        "n_with_recorded_sha256": sum(1 for d in rec.get("distributions", []) if d.get("sha256")),
        "n_match": verdicts.count("match"),
        "n_mismatch": verdicts.count("mismatch") + verdicts.count("mismatch-derived"),
        "n_mismatch_live": verdicts.count("mismatch-live"),
        "n_blocked": verdicts.count("blocked"),
        "n_not_reproducible": verdicts.count("not-reproducible"),
        "n_http_error": verdicts.count("http-error") + verdicts.count("fetch-error"),
        "n_no_checksum": verdicts.count("no-checksum-recorded"),
    }
    out["file_level_recorded"] = file_level_recorded(rec)
    out["figshare_articles"] = figshare_articles(rec)

    # ---- table re-derivation
    try:
        if rid.startswith("amniote"):
            csv_url = next(d["download_url"] for d in rec["distributions"] if d["download_url"].endswith(".csv"))
            b = cached_bytes(csv_url)
            if b:
                df, meta = parse_amniote(b)
                out["table"] = meta["info"]
                out["taxa_rechecks"] = compare_taxa(rec, df, meta["cols"], taxa)
            zips = [d["download_url"] for d in rec["distributions"] if "figshare" in d["download_url"]]
            for zu in zips:
                zb = cached_bytes(zu)
                if zb:
                    try:
                        z = zipfile.ZipFile(io.BytesIO(zb))
                        members = [(i.filename, i.file_size) for i in z.infolist()]
                        inner = [m for m in members if m[0].lower().endswith(".csv")]
                        inner_rows = None
                        inner_sha = None
                        if inner:
                            ib = z.read(inner[0][0])
                            inner_sha = sha256_bytes(ib)
                            inner_rows = int(len(pd.read_csv(io.BytesIO(ib), low_memory=False)))
                        out["figshare_zip"] = {"url": zu, "members": members[:20], "inner_csv_rows": inner_rows,
                                               "inner_csv_sha256": inner_sha,
                                               "inner_csv_equals_esapubs_csv": (inner_sha == sha256_bytes(b)) if b and inner_sha else None}
                    except zipfile.BadZipFile:
                        out["figshare_zip"] = {"url": zu, "error": "not a zip"}
        elif rid.startswith("anage"):
            b = cached_bytes("https://genomics.senescence.info/species/dataset.zip")
            if b:
                df, meta = parse_anage_zip(b)
                out["table"] = meta["info"]
                out["taxa_rechecks"] = compare_taxa(rec, df, meta["cols"], taxa)
        elif rid.startswith("pantheria"):
            b = cached_bytes("https://esapubs.org/archive/ecol/E090/184/PanTHERIA_1-0_WR05_Aug2008.txt")
            if b:
                df, meta = parse_pantheria(b)
                out["table"] = meta["info"]
                out["taxa_rechecks"] = compare_taxa(rec, df, meta["cols"], taxa)
        elif rid.startswith("eltontraits"):
            bb = cached_bytes("https://ndownloader.figshare.com/files/5631081") or cached_bytes("https://esapubs.org/archive/ecol/E095/178/BirdFuncDat.txt")
            mb = cached_bytes("https://ndownloader.figshare.com/files/5631084") or cached_bytes("https://esapubs.org/archive/ecol/E095/178/MamFuncDat.txt")
            if bb and mb:
                bdf, bmeta = parse_elton(bb, "bird")
                mdf, mmeta = parse_elton(mb, "mammal")
                out["table"] = {"birds": bmeta["info"], "mammals": mmeta["info"]}
                # taxa: compare against the union (species-rank by binomial in either file)
                rb = compare_taxa(rec, bdf, bmeta["cols"], taxa)
                rm = compare_taxa(rec, mdf, mmeta["cols"], taxa)
                # merge: a taxon matches if either file reproduces the recorded count, or the sum does
                recorded = recorded_per_taxon(rec)
                comps = []
                for slug, n in recorded.items():
                    name, rank = taxon_name_rank(slug, taxa)
                    if not name:
                        continue
                    ob = count_matches(bdf, bmeta["cols"], name, rank) or 0
                    om = count_matches(mdf, mmeta["cols"], name, rank) or 0
                    comps.append({"taxon": slug, "rank": rank, "recorded": n, "observed_birds": ob,
                                  "observed_mammals": om, "match": n in (ob, om, ob + om)})
                out["taxa_rechecks"] = {"n_recorded_per_taxon": len(recorded),
                                        "n_compared": len(comps),
                                        "n_match": sum(1 for c in comps if c["match"]),
                                        "mismatches": [c for c in comps if not c["match"]][:60],
                                        "unmatched_species_rank_taxa": sorted(set(rb["unmatched_species_rank_taxa"]) & set(rm["unmatched_species_rank_taxa"]))}
        elif rid.startswith("avonet"):
            b = cached_bytes("https://ndownloader.figshare.com/files/34480856")
            if b:
                df, meta = parse_avonet(b)
                out["table"] = meta["info"]
                out["taxa_rechecks"] = compare_taxa(rec, df, meta["cols"], taxa)
                out["file_md5"] = hashlib.md5(b).hexdigest()
        elif rid.startswith("fishbase-v25"):
            sb = cached_bytes("https://huggingface.co/datasets/cboettig/fishbase/resolve/main/data/fb/v25.04/parquet/species.parquet")
            fb = cached_bytes("https://huggingface.co/datasets/cboettig/fishbase/resolve/main/data/fb/v25.04/parquet/families.parquet")
            if sb and fb:
                df, meta = parse_fishbase_parquet(sb, fb)
                out["table"] = meta["info"]
                out["taxa_rechecks"] = compare_taxa(rec, df, meta["cols"], taxa)
        elif rid.startswith("fishbase"):
            out["api_rechecks"] = recheck_fishbase_pages(rec, dist_results)
        elif rid.startswith("morphosource"):
            out["api_rechecks"] = recheck_morphosource(rec, dist_results, taxa)
        elif rid.startswith("worms"):
            out["api_rechecks"] = recheck_worms(rec, dist_results)
        elif rid.startswith("wikidata"):
            out["api_rechecks"] = recheck_wikidata(rec, dist_results, taxa)
        elif rid.startswith("open") or rid.startswith("opentree"):
            out["api_rechecks"] = recheck_opentree(rec, dist_results, taxa)
        elif rid.startswith("timetree"):
            out["api_rechecks"] = recheck_timetree(rec, dist_results)
        elif rid.startswith("iucn"):
            out["api_rechecks"] = {"note": "all IUCN URLs answered HTTP 403 (anti-bot); ids could not be re-resolved",
                                   "statuses": sorted({d.get("status") for d in dist_results}, key=str)}
    except Exception as e:  # noqa: BLE001
        out["recheck_error"] = f"{type(e).__name__}: {e}"
        print(f"  ! {rid}: {out['recheck_error']}", file=sys.stderr)

    # ---- whole-file count comparisons
    tbl = out.get("table") or {}
    fl = out["file_level_recorded"]
    cmp: list[dict] = []

    def add(pid, observed):
        for r in fl.get(pid, []):
            cmp.append({"property_id": pid, "slice": r["slice"], **within_tolerance(r["value"], observed)})

    if rid.startswith(("amniote", "anage", "pantheria")):
        add("n_species", tbl.get("n_rows"))
        add("n_rows", tbl.get("n_rows"))
        add("n_columns", tbl.get("n_columns"))
    elif rid.startswith("eltontraits") and tbl:
        add("n_bird_species", tbl["birds"]["n_rows_with_scientific_name"])
        add("n_mammal_species", tbl["mammals"]["n_rows_with_msw3_id"])
        add("n_species", (tbl["birds"]["n_rows_with_scientific_name"] or 0) + (tbl["mammals"]["n_rows_with_msw3_id"] or 0))
        add("n_rows", None)
    elif rid.startswith("avonet") and tbl:
        add("n_species", tbl.get("n_rows_avonet1_birdlife"))
        add("n_rows", tbl.get("n_rows_avonet1_birdlife"))
        add("n_species_avonet1_birdlife", tbl.get("n_rows_avonet1_birdlife"))
        add("n_columns", tbl.get("n_columns"))
    elif rid.startswith("fishbase-v25") and tbl:
        add("n_species", tbl.get("n_species_rows"))
        add("n_rows", tbl.get("n_species_rows"))
        add("n_families", tbl.get("n_family_rows"))
    out["file_count_rechecks"] = cmp
    return out


def main() -> None:
    global OFFLINE
    ap = argparse.ArgumentParser()
    ap.add_argument("--ids", help="comma-separated subset of dataset ids")
    ap.add_argument("--offline", action="store_true", help="use only cached responses")
    ap.add_argument("--out", default=str(OUT))
    args = ap.parse_args()
    OFFLINE = args.offline

    records = json.loads(MERGED.read_text())
    batch = [r for r in records if not r["id"].startswith(EXCLUDED_PREFIXES)]
    if args.ids:
        want = set(args.ids.split(","))
        batch = [r for r in batch if r["id"] in want]
    taxa = load_taxa()

    outp = Path(args.out)
    existing: dict[str, dict] = {}
    if outp.exists():
        for r in json.loads(outp.read_text()):
            existing[r["dataset_id"]] = r
    for rec in batch:
        existing[rec["id"]] = recheck_record(rec, taxa)
        order = [r["id"] for r in records if not r["id"].startswith(EXCLUDED_PREFIXES)]
        outp.write_text(json.dumps([existing[i] for i in order if i in existing], indent=1, ensure_ascii=False) + "\n")
    print(f"wrote {outp} ({len(existing)} records)", file=sys.stderr)


if __name__ == "__main__":
    main()
